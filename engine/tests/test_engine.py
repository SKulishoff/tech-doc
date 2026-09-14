from pathlib import Path

from openpyxl import load_workbook

from techdoc_engine.core import (
    ComponentRecord, IntervalRule, MaintenanceOperationCandidate, MatrixMarks,
    ReviewStatus, SourceLocator, TableKind, canonical_is, choose_candidate_pages,
    classify_table, detect_matrix_columns, extract_periodicity_rows, marks_from_row,
    match_component, pages_to_ranges, parse_interval, score_page,
)
from techdoc_engine.pipeline import secure_temp_copy
from techdoc_engine.storage import PLAN_HEADERS, connect, export_plan_xlsx, insert_operation, register_document


def test_fast_scan_and_ranges():
    s=score_page("4 Техническое обслуживание. Периодичность IS100 12 500 км",4)
    assert s.score>=15 and "maintenance" in s.reasons and "is-code" in s.reasons
    _,pages=choose_candidate_pages(["Описание","техническое обслуживание IS100 периодичность","Упаковка"],5,1)
    assert pages==[1,2,3]
    assert pages_to_ranges([1,2,3,7,8])==[(1,3),(7,8)]


def test_interval_parser():
    assert canonical_is("IS 540") == "IS540"
    r=parse_interval("12 500 км ± 20 %"); assert r and r.value==12500 and r.unit=="km" and r.tolerance_percent==20
    r=parse_interval("не более 30 календарных дней"); assert r and r.value==30 and r.unit=="day" and r.qualifier=="max"
    r=parse_interval("не менее 1 раза в 10 дней"); assert r and r.value==10 and r.unit=="day" and r.qualifier=="min"
    r=parse_interval("не позднее 8 лет эксплуатации"); assert r and r.value==8 and r.unit=="year" and r.qualifier=="max"


def test_periodicity_and_matrix_tables():
    rows=[["Вид ТОиР","Межремонтный пробег, км"],["IS100 (осмотр)","12 500 км ± 20 %"],["IS540","600 000 км ± 20 %"]]
    assert classify_table(rows)==TableKind.PERIODICITY
    got={k:v.value for k,v in extract_periodicity_rows(rows).items()}; assert got=={"IS100":12500,"IS540":600000}
    header=[["Элемент","Работа по техническому обслуживанию","IS100","IS200","IS540","IS700"]]
    cols=detect_matrix_columns(header); assert cols=={"IS100":2,"IS200":3,"IS540":4,"IS700":5}
    marks=marks_from_row(["Фильтр","Проверить","+","-","Х",""],cols)
    assert marks.IS100 and marks.IS540 and not marks.IS200 and not marks.IS700


def test_component_matching_is_conservative():
    comps=[ComponentRecord("GC-20-10","Механическая часть головной сцепки"),ComponentRecord("CB-20-10","Компрессорный агрегат")]
    exact=match_component("Компрессорный агрегат",comps); assert exact.code=="CB-20-10" and not exact.needs_review
    alias=match_component("АКВ компрессор",comps,{"акв компрессор":"CB-20-10"}); assert alias.code=="CB-20-10" and not alias.needs_review
    weak=match_component("совершенно другой узел",comps); assert weak.needs_review


def test_duplicate_storage_and_exact_export(tmp_path: Path):
    con=connect(tmp_path/"db.sqlite")
    did,status=register_document(con,code="TEST РЭ",sha256="a"*64,filename="test.pdf"); assert status=="new"
    did2,status2=register_document(con,code="TEST РЭ",sha256="a"*64,filename="test.pdf"); assert did2==did and status2=="duplicate"
    _,status3=register_document(con,code="TEST РЭ",sha256="b"*64,filename="test_v2.pdf"); assert status3=="revision"
    con.execute("INSERT INTO components(code,subgroup,name) VALUES('CB-20-10','Компрессор','Компрессорный агрегат')"); con.commit()
    op=MaintenanceOperationCandidate(component_code="CB-20-10",operation_code="CB-20-10-01",operation_name="Проверка фильтра",operation_type="Контроль и проверки",
       interval=IntervalRule(verbatim="300 000 км",value=300000,unit="km",qualifier="exact"),marks=MatrixMarks(IS530=True,IS540=True),
       regulating_document="TEST РЭ",source=SourceLocator(document_code="TEST РЭ",page=23),confidence=1,review_status=ReviewStatus.APPROVED)
    insert_operation(con,did,op); out=export_plan_xlsx(con,tmp_path/"plan.xlsx")
    ws=load_workbook(out,read_only=True)["План ТО"]
    assert [c.value for c in next(ws.iter_rows(max_row=1))]==PLAN_HEADERS
    row=[c.value for c in next(ws.iter_rows(min_row=2,max_row=2))]
    assert row[PLAN_HEADERS.index("IS530")]=="+" and row[PLAN_HEADERS.index("IS540")]=="+" and row[PLAN_HEADERS.index("IS100")]=="-"


def test_temp_source_is_deleted(tmp_path: Path):
    src=tmp_path/"source.pdf"; src.write_bytes(b"test")
    with secure_temp_copy(src) as temp:
        assert temp.exists(); temp_path=temp
    assert not temp_path.exists()
